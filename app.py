from pathlib import Path
from datetime import date, timedelta, datetime
import io
import re
import zipfile
import pandas as pd
import streamlit as st
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font

from render_pitch_map import generate_pitch_map, get_cached_config_dfs
from render_front_cover import generate_front_cover
from render_locnotes import generate_locnotes_page
from render_pitch_change import generate_pitch_change_map
from render_eventlogs import generate_eventlogs_page
from render_coc_partners import generate_chairwelcome_page, generate_coc_page, generate_partners_page
from render_pdf import compile_fixture_pdf

BASE_DIR = Path(__file__).resolve().parent
CONFIG_PATH = BASE_DIR / "config.xlsx"
ASSETS_DIR = BASE_DIR / "assets"
OPPOSITIONS_DIR = ASSETS_DIR / "oppositions"
OUTPUT_DIR = BASE_DIR / "output"

DEFAULT_NOTES = (
    "Please note the club is expected to be very busy this weekend. "
    "We kindly ask all visiting and home families to car-share wherever possible "
    "and follow the directions of our parking marshals."
)

st.set_page_config(
    page_title="HRFC Match-Day Asset Generator",
    page_icon="🏉",
    layout="wide",
)


def get_next_sunday() -> date:
    today = date.today()
    days_ahead = 6 - today.weekday()
    if days_ahead <= 0:
        days_ahead += 7
    return today + timedelta(days=days_ahead)


def get_config_mtime() -> float:
    return CONFIG_PATH.stat().st_mtime if CONFIG_PATH.exists() else 0.0


def calculate_in_time(ko_str: str) -> str:
    try:
        dt = datetime.strptime(ko_str.strip(), "%H:%M")
        in_dt = dt - timedelta(minutes=75)
        return in_dt.strftime("%H:%M")
    except Exception:
        return "08:45"


def calculate_out_time(ko_str: str, match_length_mins: int) -> str:
    try:
        dt = datetime.strptime(ko_str.strip(), "%H:%M")
        total_duration = match_length_mins + 20 + 20
        out_dt = dt + timedelta(minutes=total_duration)
        return out_dt.strftime("%H:%M")
    except Exception:
        return "12:30"


def requires_changing_rooms(home_team: str) -> bool:
    clean = str(home_team).strip().upper()
    match = re.search(r"U(\d+)", clean)
    if match:
        try:
            age_num = int(match.group(1))
            if age_num < 12:
                return False
        except ValueError:
            pass
    return True


@st.cache_data
def load_base_metadata(mtime: float):
    anchors_df, _ = get_cached_config_dfs(CONFIG_PATH)

    pitch_keys = []
    if not anchors_df.empty:
        key_col = next((c for c in anchors_df.columns if str(c).strip().lower() == "pitch_key"), None)
        if key_col:
            pitch_keys = [
                str(x).strip()
                for x in anchors_df[key_col].dropna().unique().tolist()
                if str(x).strip() and str(x).strip().upper() != "NAN"
            ]
        else:
            st.error("Sheet 'pitch_anchors' was found, but column 'pitch_key' is missing.")
    else:
        st.error(f"Could not load 'pitch_anchors' sheet from {CONFIG_PATH.name}.")

    opponents_df = pd.DataFrame()
    opponents = []
    try:
        opponents_df = pd.read_excel(CONFIG_PATH, sheet_name="opponents")
        name_col = next(
            (
                c
                for c in opponents_df.columns
                if str(c).strip().lower() in ["display_name", "displayname", "display", "opponent", "club", "team"]
            ),
            None,
        )
        if name_col:
            opponents = [
                str(x).strip().upper()
                for x in opponents_df[name_col].dropna().unique().tolist()
                if str(x).strip() and str(x).strip().upper() != "NAN"
            ]
    except Exception as e:
        st.warning(f"Could not read opponents sheet: {e}")

    if not opponents and OPPOSITIONS_DIR.exists():
        for file in sorted(OPPOSITIONS_DIR.iterdir()):
            if file.suffix.lower() in [".png", ".jpg", ".jpeg"]:
                stem = re.sub(r"_WHITE$", "", file.stem, flags=re.IGNORECASE).strip().upper()
                if stem and stem not in opponents:
                    opponents.append(stem)

    opponents = sorted(list(set(opponents)))

    comps_df = pd.DataFrame()
    raw_teams = []
    try:
        comps_df = pd.read_excel(CONFIG_PATH, sheet_name="comps")
        if "team_key" in comps_df.columns:
            raw_teams = comps_df["team_key"].dropna().astype(str).str.strip().unique().tolist()
    except Exception as e:
        st.warning(f"Could not read comps sheet: {e}")

    if not raw_teams:
        try:
            teams_df = pd.read_excel(CONFIG_PATH, sheet_name="teams")
            team_col = next((c for c in teams_df.columns if 'team' in str(c).lower() or 'key' in str(c).lower()), None)
            if team_col:
                raw_teams = teams_df[team_col].dropna().astype(str).str.strip().unique().tolist()
        except Exception:
            pass

    ordered_teams = [
        "---JUNIOR BOYS--------------------------------",
        "U13", "U14", "HURRICANES", "COLTS",
        "---JUNIOR GIRLS-------------------------------",
        "WARRIORS U12", "WARRIORS U14", "WARRIORS U16",
        "---MINIS--------------------------------------",
        "U12", "U11", "U10", "U9", "U8", "U7", "U6"
    ]

    for t in raw_teams:
        if t not in ordered_teams:
            ordered_teams.append(t)

    return pitch_keys, opponents, opponents_df, comps_df, ordered_teams


def get_competitions_for_team(team_key: str, comps_df: pd.DataFrame) -> tuple[list[str], dict[str, any]]:
    options = ["NONE"]
    alias_to_info_map = {"NONE": {"code": None, "match_length": 70}}

    if comps_df.empty or "team_key" not in comps_df.columns:
        return options, alias_to_info_map

    matched = comps_df[comps_df["team_key"].astype(str).str.strip() == str(team_key).strip()]

    for _, row in matched.iterrows():
        comp_alias_raw = row.get("comp_alias")
        comp_name_raw = row.get("comp_name")
        logo_id_raw = row.get("logo_id")

        if pd.notna(comp_alias_raw):
            alias_upper = str(comp_alias_raw).strip().upper()
            if alias_upper and alias_upper not in ["NONE", "NAN", "FRIENDLY"]:
                if pd.notna(logo_id_raw) and str(logo_id_raw).strip():
                    code_str = str(logo_id_raw).strip()
                elif pd.notna(comp_name_raw) and str(comp_name_raw).strip():
                    code_str = str(comp_name_raw).strip()
                else:
                    code_str = alias_upper

                if alias_upper not in options:
                    options.append(alias_upper)
                    alias_to_info_map[alias_upper] = {"code": code_str}

    return options, alias_to_info_map


def get_default_match_length_for_team(team_key: str, comps_df: pd.DataFrame) -> int:
    if comps_df.empty or "team_key" not in comps_df.columns:
        return 70
    matched = comps_df[comps_df["team_key"].astype(str).str.strip() == str(team_key).strip()]
    if not matched.empty:
        val = matched.iloc[0].get("match_length", 70)
        try:
            return int(val)
        except (ValueError, TypeError):
            return 70
    return 70


def generate_template_excel(config_excel_path: Path) -> bytes:
    output = io.BytesIO()
    
    teams_list = ["WARRIORS U16", "HURRICANES U14"]
    opponents_list = ["ABBEY RFC", "NEWBURY RFC"]
    pitch_keys_list = ["P1", "P2", "P3", "P4"]
    comps_list = ["BERKS YOUTH CUP", "FRIENDLY", "LEAGUE"]
    rooms_list = [1, 2, 3, 4]
    
    try:
        if config_excel_path.exists():
            comps_df = pd.read_excel(config_excel_path, sheet_name="comps")
            if "team_key" in comps_df.columns:
                teams_list = comps_df["team_key"].dropna().astype(str).str.strip().unique().tolist()
            else:
                teams_df = pd.read_excel(config_excel_path, sheet_name="teams")
                team_col = next((c for c in teams_df.columns if 'team' in str(c).lower() or 'key' in str(c).lower()), None)
                if team_col:
                    teams_list = teams_df[team_col].dropna().astype(str).str.strip().unique().tolist()
                
            opps_df = pd.read_excel(config_excel_path, sheet_name="opponents")
            opp_col = next((c for c in opps_df.columns if 'display' in str(c).lower() or 'opponent' in str(c).lower()), None)
            if opp_col:
                opponents_list = opps_df[opp_col].dropna().astype(str).str.strip().unique().tolist()
                
            anchors_df = pd.read_excel(config_excel_path, sheet_name="pitch_anchors")
            pitch_col = next((c for c in anchors_df.columns if 'pitch_key' in str(c).lower()), None)
            if pitch_col:
                pitch_keys_list = anchors_df[pitch_col].dropna().astype(str).str.strip().unique().tolist()
                
            comp_col = next((c for c in comps_df.columns if 'alias' in str(c).lower() or 'name' in str(c).lower()), None)
            if comp_col:
                comps_list = comps_df[comp_col].dropna().astype(str).str.strip().unique().tolist()
    except Exception:
        pass

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_template = pd.DataFrame(
            {
                "home_team": [teams_list[0] if teams_list else "WARRIORS U16"],
                "opponent": [opponents_list[0] if opponents_list else "ABBEY RFC"],
                "pitch_key": [pitch_keys_list[0] if pitch_keys_list else "P1"],
                "ko_time": ["10:00"],
                "match_date": ["13 SEP 2026"],
                "referee": ["TBC"],
                "competition": [comps_list[0] if comps_list else "FRIENDLY"],
                "is_provisional": ["FALSE"],
                "home_room_num": [1],
                "away_room_num": [3],
                "custom_notes": [
                    "Please park in the main car park."
                ],
            }
        )
        df_template.to_excel(writer, sheet_name="Fixtures", index=False)
        
        workbook = writer.book
        lookup_sheet = workbook.create_sheet(title="Lists")
        
        for r_idx, val in enumerate(teams_list, start=1):
            lookup_sheet.cell(row=r_idx, column=1, value=val)
        for r_idx, val in enumerate(opponents_list, start=1):
            lookup_sheet.cell(row=r_idx, column=2, value=val)
        for r_idx, val in enumerate(pitch_keys_list, start=1):
            lookup_sheet.cell(row=r_idx, column=3, value=val)
        for r_idx, val in enumerate(comps_list, start=1):
            lookup_sheet.cell(row=r_idx, column=4, value=val)
        for r_idx, val in enumerate(rooms_list, start=1):
            lookup_sheet.cell(row=r_idx, column=5, value=val)
        for r_idx, val in enumerate(["TRUE", "FALSE"], start=1):
            lookup_sheet.cell(row=r_idx, column=6, value=val)

        font_header = Font(name="Aptos Narrow", size=9, bold=True)
        font_body = Font(name="Aptos Narrow", size=9, bold=False)

        for sheet in workbook.worksheets:
            sheet.views.sheetView[0].showGridLines = False
            for col in range(1, sheet.max_column + 1):
                col_letter = sheet.cell(row=1, column=col).column_letter
                max_len = 0
                for row in range(1, sheet.max_row + 1):
                    cell = sheet.cell(row=row, column=col)
                    if row == 1:
                        cell.font = font_header
                    else:
                        cell.font = font_body
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
                sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

        ws = workbook["Fixtures"]
        max_row = 100
        
        def add_list_validation(col_letter, max_r, formula):
            dv = DataValidation(type="list", formula1=formula, allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col_letter}2:{col_letter}{max_r}")

        add_list_validation("A", max_row, f"=Lists!$A$1:$A${len(teams_list)}")
        add_list_validation("B", max_row, f"=Lists!$B$1:$B${len(opponents_list)}")
        add_list_validation("C", max_row, f"=Lists!$C$1:$C${len(pitch_keys_list)}")
        add_list_validation("G", max_row, f"=Lists!$D$1:$D${len(comps_list)}")
        add_list_validation("H", max_row, "=Lists!$F$1:$F$2")
        add_list_validation("I", max_row, f"=Lists!$E$1:$E${len(rooms_list)}")
        add_list_validation("J", max_row, f"=Lists!$E$1:$E${len(rooms_list)}")

    return output.getvalue()


pitch_keys, opponent_list, opponents_meta_df, comps_df, home_team_options = load_base_metadata(
    get_config_mtime()
)

st.title("🏉 HRFC Match-Day Asset Generator")

mode = st.radio(
    "Mode",
    ["Single Fixture", "Batch Processing"],
    horizontal=True,
    label_visibility="collapsed",
)

if mode == "Single Fixture":
    col_params, col_preview = st.columns([1, 1], gap="large")

    with col_params:
        st.subheader("Fixture Parameters")

        selected_pitch = st.selectbox(
            "Select Pitch Allocation",
            options=pitch_keys,
            index=0 if pitch_keys else None,
        )

        home_team = st.selectbox(
            "Home Team",
            options=home_team_options if home_team_options else ["WARRIORS U16"],
            index=0,
        )

        opponent = st.selectbox(
            "Opponent",
            options=opponent_list if opponent_list else ["ABBEY RFC"],
            index=0,
        )

        col_ko, col_date = st.columns(2)
        with col_ko:
            ko_time = st.text_input("Kick-Off Time", value="10:00")
        with col_date:
            match_date_val = st.date_input(
                "Match Date",
                value=get_next_sunday(),
                format="DD/MM/YYYY",
            )
            match_date_str = match_date_val.strftime("%d %b %Y").upper()
            st.caption(f"Formatted: **{match_date_str}**")

        col_ref, col_comp = st.columns(2)
        with col_ref:
            referee = st.text_input("Referee", value="TBC")

        with col_comp:
            comp_options, alias_to_info_map = get_competitions_for_team(home_team, comps_df)
            selected_alias = st.selectbox(
                "Competition",
                options=comp_options,
                index=0,
                key=f"comp_select_{home_team}"
            )
            comp_info = alias_to_info_map.get(selected_alias, {"code": None})
            comp_code = comp_info["code"]

        match_length = get_default_match_length_for_team(home_team, comps_df)

        is_provisional = st.checkbox("Mark as PROVISIONAL", value=False)

        matchday_notes = st.text_area(
            "Club / Parking Commentary (Page 5 Notes)",
            value=DEFAULT_NOTES,
            height=90,
            help="Dynamic text displayed to the right of the QR code on the location notes plate.",
        )

        show_changing_rooms = requires_changing_rooms(home_team)

        home_room_num = 1
        away_room_num = 3
        in_time = calculate_in_time(ko_time)
        out_time = calculate_out_time(ko_time, match_length)

        if show_changing_rooms:
            st.markdown("---")
            st.subheader("Changing Room Settings")
            
            all_rooms = [1, 2, 3, 4]
            col_r1, col_r2 = st.columns(2)
            with col_r1:
                home_room_num = st.selectbox(home_team, all_rooms, index=0, key="home_room")
            
            away_room_options = [r for r in all_rooms if r != home_room_num]
            with col_r2:
                current_away_idx = 0
                if away_room_num in away_room_options:
                    current_away_idx = away_room_options.index(away_room_num)
                away_room_num = st.selectbox(opponent, away_room_options, index=current_away_idx, key="away_room")

            col_in, col_out = st.columns(2)
            with col_in:
                st.metric(label="IN", value=in_time)
            with col_out:
                st.metric(label="OUT", value=out_time)

        opp_alias = None
        opp_crest = None
        if not opponents_meta_df.empty:
            lookup_col = next(
                (
                    c
                    for c in opponents_meta_df.columns
                    if str(c).strip().lower() in ["display_name", "displayname", "display", "opponent"]
                ),
                None,
            )
            if lookup_col:
                matched_opp = opponents_meta_df[
                    opponents_meta_df[lookup_col].astype(str).str.strip().str.upper() == str(opponent).strip().upper()
                ]
                if not matched_opp.empty:
                    row = matched_opp.iloc[0]
                    alias_col = next((c for c in opponents_meta_df.columns if "alias" in str(c).lower()), None)
                    crest_col = next((c for c in opponents_meta_df.columns if "crest" in str(c).lower()), None)

                    if alias_col and pd.notna(row[alias_col]):
                        opp_alias = str(row[alias_col]).strip()
                    if crest_col and pd.notna(row[crest_col]):
                        opp_crest = str(row[crest_col]).strip()

        generate_btn = st.button("Generate Assets", type="primary", use_container_width=True)

    with col_preview:
        st.subheader("Asset Preview")

        if generate_btn:
            if "---" in str(home_team):
                st.error("Please select a valid Home Team from the list (section dividers cannot be used as a team).")
            else:
                with st.spinner("Generating graphics & PDF..."):
                    cover_path = generate_front_cover(
                        home_team=home_team,
                        opponent=opponent,
                        ko_time=ko_time,
                        match_date=match_date_str,
                        referee=referee,
                        opponent_crest_stem=opp_crest,
                        competition=comp_code,
                        output_filename="preview_front_cover.png",
                    )

                    chairwelcome_path = generate_chairwelcome_page(
                        home_team=home_team,
                        output_filename="preview_chairwelcome.png",
                    )

                    pitch_path = generate_pitch_map(
                        config_excel_path=CONFIG_PATH,
                        pitch_key=selected_pitch,
                        home_team=home_team,
                        opponent=opponent,
                        ko_time=ko_time,
                        opponent_alias=opp_alias,
                        opponent_crest_stem=opp_crest,
                        match_date=match_date_str,
                        is_provisional=is_provisional,
                        competition=comp_code,
                        output_filename="preview_pitch_map.png",
                    )

                    locnotes_path = generate_locnotes_page(
                        home_team=home_team,
                        opponent=opponent,
                        opponent_crest_stem=opp_crest,
                        competition=comp_code,
                        custom_notes=matchday_notes,
                        output_filename="preview_locnotes.png",
                    )

                    eventlogs_path = generate_eventlogs_page(
                        config_excel_path=CONFIG_PATH,
                        home_team=home_team,
                        ko_time=ko_time,
                        match_length_mins=match_length,
                        output_filename="preview_eventlogs.png",
                    )

                    coc_path = generate_coc_page(
                        home_team=home_team,
                        output_filename="preview_coc.png",
                    )

                    partners_path = generate_partners_page(
                        home_team=home_team,
                        output_filename="preview_partners.png",
                    )

                    st.session_state["cover_img"] = cover_path
                    st.session_state["chairwelcome_img"] = chairwelcome_path
                    st.session_state["pitch_img"] = pitch_path
                    st.session_state["locnotes_img"] = locnotes_path
                    st.session_state["eventlogs_img"] = eventlogs_path
                    st.session_state["coc_img"] = coc_path
                    st.session_state["partners_img"] = partners_path

                    ordered_pages = [cover_path, chairwelcome_path, locnotes_path]

                    if show_changing_rooms:
                        pitch_change_path = generate_pitch_change_map(
                            config_excel_path=CONFIG_PATH,
                            home_team=home_team,
                            opponent=opponent,
                            home_room_num=home_room_num,
                            away_room_num=away_room_num,
                            in_time=in_time,
                            out_time=out_time,
                            category_title=f"{home_team} CHANGING ROOMS".upper(),
                            home_crest_stem="WARRIORS_CRESTDARK" if "WARRIOR" in str(home_team).upper() else None,
                            away_crest_stem=opp_crest,
                            pitch_map_path=pitch_path,
                            output_filename="preview_pitch_change.png",
                        )
                        st.session_state["pitch_change_img"] = pitch_change_path
                        ordered_pages.append(pitch_change_path)
                    else:
                        if "pitch_change_img" in st.session_state:
                            del st.session_state["pitch_change_img"]
                        ordered_pages.append(pitch_path)

                    ordered_pages.extend([eventlogs_path, coc_path, partners_path])

                    pdf_path = compile_fixture_pdf(
                        image_paths=ordered_pages,
                        output_filename=f"fixture_{home_team}_vs_{opponent}.pdf".replace(" ", "_")
                    )
                    st.session_state["fixture_pdf"] = pdf_path

        if "fixture_pdf" in st.session_state and Path(st.session_state["fixture_pdf"]).exists():
            with open(st.session_state["fixture_pdf"], "rb") as f:
                st.download_button(
                    label="📥 Download Complete Matchday PDF",
                    data=f,
                    file_name=Path(st.session_state["fixture_pdf"]).name,
                    mime="application/pdf",
                    use_container_width=True,
                )
            st.markdown("---")

        tabs_list = ["Front Cover", "Chair's Welcome", "Pitch Map", "Location & Parking"]
        if show_changing_rooms:
            tabs_list.append("Pitch Change")
        tabs_list.extend(["Event Logs", "Code of Conduct", "Partners"])

        tabs = st.tabs(tabs_list)

        with tabs[0]:
            if "cover_img" in st.session_state and Path(st.session_state["cover_img"]).exists():
                st.image(str(st.session_state["cover_img"]), use_container_width=True)
                with open(st.session_state["cover_img"], "rb") as f:
                    st.download_button(
                        label="Download Cover (PNG)",
                        data=f,
                        file_name=f"cover_{home_team}_v_{opponent}.png".replace(" ", "_"),
                        mime="image/png",
                        use_container_width=True,
                    )
            else:
                st.info("Click 'Generate Assets' to preview.")

        with tabs[1]:
            if "chairwelcome_img" in st.session_state and Path(st.session_state["chairwelcome_img"]).exists():
                st.image(str(st.session_state["chairwelcome_img"]), use_container_width=True)
                with open(st.session_state["chairwelcome_img"], "rb") as f:
                    st.download_button(
                        label="Download Chair's Welcome (PNG)",
                        data=f,
                        file_name=f"chairwelcome_{home_team}.png".replace(" ", "_"),
                        mime="image/png",
                        use_container_width=True,
                    )
            else:
                st.info("Click 'Generate Assets' to preview.")

        with tabs[2]:
            if "pitch_img" in st.session_state and Path(st.session_state["pitch_img"]).exists():
                st.image(str(st.session_state["pitch_img"]), use_container_width=True)
                with open(st.session_state["pitch_img"], "rb") as f:
                    st.download_button(
                        label="Download Pitch Map (PNG)",
                        data=f,
                        file_name=f"pitch_map_{home_team}_{selected_pitch}.png".replace(" ", "_"),
                        mime="image/png",
                        use_container_width=True,
                    )
            else:
                st.info("Click 'Generate Assets' to preview.")

        with tabs[3]:
            if "locnotes_img" in st.session_state and Path(st.session_state["locnotes_img"]).exists():
                st.image(str(st.session_state["locnotes_img"]), use_container_width=True)
                with open(st.session_state["locnotes_img"], "rb") as f:
                    st.download_button(
                        label="Download Location & Parking (PNG)",
                        data=f,
                        file_name=f"locnotes_{home_team}_v_{opponent}.png".replace(" ", "_"),
                        mime="image/png",
                        use_container_width=True,
                    )
            else:
                st.info("Click 'Generate Assets' to preview.")

        tab_idx = 4
        if show_changing_rooms and len(tabs) > tab_idx:
            with tabs[tab_idx]:
                if "pitch_change_img" in st.session_state and Path(st.session_state["pitch_change_img"]).exists():
                    st.image(str(st.session_state["pitch_change_img"]), use_container_width=True)
                    with open(st.session_state["pitch_change_img"], "rb") as f:
                        st.download_button(
                            label="Download Pitch Change Graphic (PNG)",
                            data=f,
                            file_name=f"pitch_change_{home_team}_v_{opponent}.png".replace(" ", "_"),
                            mime="image/png",
                            use_container_width=True,
                        )
                else:
                    st.info("Click 'Generate Assets' to preview.")
            tab_idx += 1

        if len(tabs) > tab_idx:
            with tabs[tab_idx]:
                if "eventlogs_img" in st.session_state and Path(st.session_state["eventlogs_img"]).exists():
                    st.image(str(st.session_state["eventlogs_img"]), use_container_width=True)
                    with open(st.session_state["eventlogs_img"], "rb") as f:
                        st.download_button(
                            label="Download Event Logs Graphic (PNG)",
                            data=f,
                            file_name=f"eventlogs_{home_team}_v_{opponent}.png".replace(" ", "_"),
                            mime="image/png",
                            use_container_width=True,
                        )
                else:
                    st.info("Click 'Generate Assets' to preview.")
            tab_idx += 1

        if len(tabs) > tab_idx:
            with tabs[tab_idx]:
                if "coc_img" in st.session_state and Path(st.session_state["coc_img"]).exists():
                    st.image(str(st.session_state["coc_img"]), use_container_width=True)
                    with open(st.session_state["coc_img"], "rb") as f:
                        st.download_button(
                            label="Download Code of Conduct (PNG)",
                            data=f,
                            file_name=f"coc_{home_team}.png".replace(" ", "_"),
                            mime="image/png",
                            use_container_width=True,
                        )
                else:
                    st.info("Click 'Generate Assets' to preview.")
            tab_idx += 1

        if len(tabs) > tab_idx:
            with tabs[tab_idx]:
                if "partners_img" in st.session_state and Path(st.session_state["partners_img"]).exists():
                    st.image(str(st.session_state["partners_img"]), use_container_width=True)
                    with open(st.session_state["partners_img"], "rb") as f:
                        st.download_button(
                            label="Download Partners Graphic (PNG)",
                            data=f,
                            file_name=f"partners_{home_team}.png".replace(" ", "_"),
                            mime="image/png",
                            use_container_width=True,
                        )
                else:
                    st.info("Click 'Generate Assets' to preview.")

else:
    st.subheader("Batch Processing")
    st.markdown("Download the batch upload template below, fill in your fixture parameters, and upload it to generate all assets in bulk.")
    
    st.download_button(
        label="Download Batch Upload Template (.xlsx)",
        data=generate_template_excel(CONFIG_PATH),
        file_name="HRFC_Batch_Fixtures_Template.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )

    st.markdown("---")
    uploaded_file = st.file_uploader("Upload Completed Fixtures Spreadsheet (.xlsx)", type=["xlsx"])

    if uploaded_file is not None:
        fixtures_df = pd.read_excel(uploaded_file)
        st.dataframe(fixtures_df.head(10), use_container_width=True)

        if st.button("Run Full Matchday Batch", type="primary"):
            zip_buffer = io.BytesIO()
            with st.spinner("Compiling graphics..."):
                with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
                    for idx, row in fixtures_df.iterrows():
                        team = str(row["home_team"])
                        opp = str(row["opponent"]).upper()
                        pkey = str(row["pitch_key"])
                        ko = str(row["ko_time"])
                        ref = str(row.get("referee", "TBC"))
                        mdate = str(row.get("match_date", get_next_sunday().strftime("%d %b %Y").upper()))
                        notes = str(row.get("custom_notes", DEFAULT_NOTES))
                        prov = str(row.get("is_provisional", "false")).strip().upper() in [
                            "TRUE",
                            "1",
                            "YES",
                        ]
                        h_room = int(row.get("home_room_num", 1)) if pd.notna(row.get("home_room_num")) else 1
                        a_room = int(row.get("away_room_num", 3)) if pd.notna(row.get("away_room_num")) else 3

                        c_stem = None
                        if not opponents_meta_df.empty:
                            lookup_col = next(
                                (
                                    c
                                    for c in opponents_meta_df.columns
                                    if str(c).strip().lower() in ["display_name", "displayname", "display", "opponent"]
                                ),
                                None,
                            )
                            if lookup_col:
                                matched_opp = opponents_meta_df[
                                    opponents_meta_df[lookup_col].astype(str).str.strip().str.upper() == opp
                                ]
                                if not matched_opp.empty:
                                    crest_col = next((c for c in opponents_meta_df.columns if "crest" in str(c).lower()), None)
                                    if crest_col and pd.notna(matched_opp.iloc[0][crest_col]):
                                        c_stem = str(matched_opp.iloc[0][crest_col]).strip()

                        if not c_stem:
                            c_stem = opp

                        raw_comp = (
                            str(row["competition"]).strip()
                            if "competition" in row and pd.notna(row["competition"])
                            else None
                        )
                        comp = (
                            None
                            if not raw_comp or raw_comp.upper() in ["NONE", "FRIENDLY", "NAN"]
                            else raw_comp
                        )

                        clean_team = team.replace(" ", "_")
                        clean_opp = opp.replace(" ", "_")

                        c_file = f"cover_{clean_team}_v_{clean_opp}_{idx}.png"
                        c_path = generate_front_cover(
                            home_team=team,
                            opponent=opp,
                            ko_time=ko,
                            match_date=mdate,
                            referee=ref,
                            opponent_crest_stem=c_stem,
                            competition=comp,
                            output_filename=c_file,
                        )
                        zip_file.write(c_path, arcname=f"front_covers/{c_file}")

                        cw_file = f"chairwelcome_{clean_team}_{idx}.png"
                        cw_path = generate_chairwelcome_page(home_team=team, output_filename=cw_file)
                        zip_file.write(cw_path, arcname=f"chair_welcome/{cw_file}")

                        p_file = f"pitch_{clean_team}_{pkey}_{idx}.png"
                        p_path = generate_pitch_map(
                            config_excel_path=CONFIG_PATH,
                            pitch_key=pkey,
                            home_team=team,
                            opponent=opp,
                            ko_time=ko,
                            opponent_crest_stem=c_stem,
                            match_date=mdate,
                            is_provisional=prov,
                            competition=comp,
                            output_filename=p_file,
                        )
                        zip_file.write(p_path, arcname=f"pitch_maps/{p_file}")

                        if requires_changing_rooms(team):
                            match_len = get_default_match_length_for_team(team, comps_df)
                            pc_file = f"pitch_change_{clean_team}_{idx}.png"
                            pc_path = generate_pitch_change_map(
                                config_excel_path=CONFIG_PATH,
                                home_team=team,
                                opponent=opp,
                                home_room_num=h_room,
                                away_room_num=a_room,
                                in_time=calculate_in_time(ko),
                                out_time=calculate_out_time(ko, match_len),
                                category_title=f"{team} CHANGING ROOMS".upper(),
                                home_crest_stem="WARRIORS_CRESTDARK" if "WARRIOR" in team.upper() else None,
                                away_crest_stem=c_stem,
                                pitch_map_path=p_path,
                                output_filename=pc_file,
                            )
                            zip_file.write(pc_path, arcname=f"pitch_changes/{pc_file}")

                        v_file = f"locnotes_{clean_team}_v_{clean_opp}_{idx}.png"
                        ln_path = generate_locnotes_page(
                            home_team=team,
                            opponent=opp,
                            opponent_crest_stem=c_stem,
                            competition=comp,
                            custom_notes=notes,
                            output_filename=v_file,
                        )
                        zip_file.write(ln_path, arcname=f"locnotes_pages/{v_file}")

                        el_file = f"eventlogs_{clean_team}_{idx}.png"
                        match_len = get_default_match_length_for_team(team, comps_df)
                        el_path = generate_eventlogs_page(
                            config_excel_path=CONFIG_PATH,
                            home_team=team,
                            ko_time=ko,
                            match_length_mins=match_len,
                            output_filename=el_file,
                        )
                        zip_file.write(el_path, arcname=f"event_logs/{el_file}")

                        coc_file = f"coc_{clean_team}_{idx}.png"
                        coc_path = generate_coc_page(home_team=team, output_filename=coc_file)
                        zip_file.write(coc_path, arcname=f"coc/{coc_file}")

                        part_file = f"partners_{clean_team}_{idx}.png"
                        part_path = generate_partners_page(home_team=team, output_filename=part_file)
                        zip_file.write(part_path, arcname=f"partners/{part_file}")

            st.success("Batch generation complete.")
            st.download_button(
                label="Download Assets (.zip)",
                data=zip_buffer.getvalue(),
                file_name="HRFC_Matchday_Assets.zip",
                mime="application/zip",
                use_container_width=True,
            )